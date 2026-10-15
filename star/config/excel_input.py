"""
The Excel input template: write it, read it back, report every problem.

Differences from the legacy template (stardust_app.py 83-420):

* Rows are matched on a stable key in column A, not on the label text. The legacy
  ignored any row whose label drifted by one character from `PARAM_KEY_MAP`.
* Enumerated settings get dropdowns, so an invalid choice cannot be typed.
* A value that cannot be read is reported with its sheet and row. The legacy's
  `except Exception: pass` silently replaced it with a default.
* Forecast horizons and mandate limits are first-class inputs.
* The template itself is set up to print on A4 landscape.
"""
from __future__ import annotations

import math
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..core.errors import ConfigError
from .schema import (DRIFTS, ENGINES, REBALANCE_FREQS, SIGNAL_MODES, VIEW_STYLES, StarConfig,
                     normalise_tickers)

#: (key, label, guidance) in display order; section banners are rows with key None
LAYOUT: tuple[tuple[str | None, str, str], ...] = (
    (None, "RUN", ""),
    ("initial_value", "Portfolio value (Rs.)", "Capital the rupee figures are based on"),
    ("as_of", "As-of date (YYYY-MM-DD)", "Blank = latest available trading day. Fix it to reproduce a past run"),
    (None, "DATA", ""),
    ("data.benchmark", "Benchmark index", "^NSEI = NIFTY 50 (price index; total return is built from NIFTYBEES)"),
    ("data.tri_proxy", "Total-return proxy", "NIFTYBEES.NS: calibrates the dividend accrual"),
    ("data.lookback_years", "Lookback (years)", "History used for estimation"),
    ("data.max_missing", "Max missing data per stock", "0.05 = exclude a stock missing over 5% of sessions (listed late, suspended)"),
    ("data.min_history_sessions", "Min history (sessions)", "Stocks with less history are excluded, with the reason shown"),
    ("data.min_adv_inr", "Min daily value traded (Rs.)", "Liquidity floor in rupees, not shares"),
    ("data.risk_free_csv", "Risk-free rate file (CSV)", "Optional. Blank = RBI 91-day T-bill yields from the bundled IIMA data"),
    ("data.iima_factor_file", "IIMA factor file", "Optional. Blank = bundled IIMA release (four factors + T-bill rate)"),
    ("data.risk_free_fallback", "Fallback risk-free rate", "0.065 = 6.5%, used only if no dated rate can be loaded (flagged)"),
    (None, "FORECAST", ""),
    ("forecast.horizons", "Horizons", "Comma-separated: 1d, 2d, 1w, 1m, 3m, 6m, 1y, or trading days such as 45"),
    ("forecast.default_horizon", "Default horizon shown", "The horizon the summary opens on"),
    ("forecast.engine", "Scenario engine", "regime_fhs = regime-aware historical simulation (recommended)"),
    ("forecast.n_scenarios", "Scenarios", "20,000 is plenty; more is slower, not more accurate"),
    ("forecast.drift", "Return assumption", "prior = equilibrium returns (recommended); views = your views are right"),
    ("forecast.seed", "Random seed", "Same seed and data = same numbers"),
    (None, "EXPECTED RETURNS", ""),
    ("returns.target_market_sharpe", "Market Sharpe ratio assumed", "0.35 implies an equity premium near 5%"),
    ("returns.tau", "Prior uncertainty (tau)", "0.05 is standard"),
    ("returns.view_style", "Views", "relative_quintile = best vs worst fifth of the signal; none = equilibrium only"),
    ("returns.view_confidence", "View confidence", "0.30 = views move weights 30% of the way to full conviction"),
    (None, "MANDATE LIMITS", ""),
    ("mandate.max_stock_weight", "Max weight per stock", "0.10 = 10%"),
    ("mandate.max_sector_weight", "Max weight per sector", "0.30 = 30%"),
    ("mandate.min_holdings", "Min holdings", ""),
    ("mandate.max_holdings", "Max holdings", "Enforced, not just reported"),
    ("mandate.max_annual_turnover", "Max annual turnover", "1.50 = 150% one-way a year"),
    ("mandate.max_drawdown_tolerance", "Drawdown tolerance", "0.20 = 20% peak-to-trough"),
    ("mandate.max_adv_participation", "Max share of daily volume", "0.20 = trade at most 20% of a day's value"),
    ("mandate.max_days_to_exit", "Max days to exit a position", ""),
    ("mandate.max_expected_shortfall_1y", "Max 1-year expected shortfall", "Optional, e.g. 0.25. Blank = no limit"),
    ("mandate.max_tracking_error", "Max tracking error", "Optional, e.g. 0.08. Blank = no limit"),
    (None, "TRADING COSTS (per leg)", ""),
    ("costs.brokerage_rate", "Brokerage rate", "0.0003 = 0.03%"),
    ("costs.brokerage_cap_inr", "Brokerage cap per order (Rs.)", ""),
    ("costs.stt_delivery", "STT (delivery)", "0.001 = 0.1%, both legs"),
    ("costs.exchange_txn", "Exchange transaction charge", "NSE"),
    ("costs.sebi_fee", "SEBI turnover fee", ""),
    ("costs.stamp_duty_buy", "Stamp duty (buy only)", "0.00015 = 0.015%"),
    ("costs.gst_rate", "GST on brokerage and fees", "0.18 = 18%"),
    ("costs.dp_charge_inr_per_scrip_sell", "DP charge per stock sold (Rs.)", ""),
    ("costs.impact_coefficient", "Market impact coefficient", "Square-root impact; 1.0 is a common starting point"),
    (None, "BACKTEST", ""),
    ("backtest.rebalance", "Rebalance frequency", "ME = monthly, 3ME = quarterly, 6ME = half-yearly, YE = yearly"),
    ("backtest.window_years", "Estimation window (years)", "Must be shorter than the lookback"),
    ("backtest.signals_mode", "Signals in backtest", "price_only is point-in-time safe; with_fundamentals is biased and labelled so"),
    ("backtest.turnover_penalty", "Turnover penalty", "Discourages small, costly trades"),
    ("backtest.no_trade_band", "No-trade band", "0.005 = skip trades smaller than 0.5% of the portfolio"),
)

_CHOICES = {
    "forecast.engine": ENGINES,
    "forecast.drift": DRIFTS,
    "returns.view_style": VIEW_STYLES,
    "backtest.signals_mode": SIGNAL_MODES,
    "backtest.rebalance": REBALANCE_FREQS,
}
_OPTIONAL = {"as_of", "data.risk_free_csv", "data.iima_factor_file", "mandate.max_expected_shortfall_1y",
             "mandate.max_tracking_error"}

_HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_SECTION_FONT = Font(name="Arial", size=9, bold=True, color="1F3A5F")
_SECTION_FILL = PatternFill("solid", fgColor="E8EEF5")
_BODY = Font(name="Arial", size=9)
_NOTE = Font(name="Arial", size=8, italic=True, color="555555")
_INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")


def _print_setup(ws, last_col: str) -> None:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"
    ws.oddFooter.center.text = "Page &P of &N"


def write_template(path: str | Path, config: StarConfig | None = None) -> Path:
    """Write an input workbook pre-filled with `config` (or the defaults)."""
    cfg = config or StarConfig(tickers=())
    values = cfg.flat()
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Settings"
    for col, (head, width) in enumerate((("Key", 30), ("Setting", 34), ("Value", 22), ("Guidance", 70)), 1):
        c = ws.cell(1, col, head)
        c.font, c.fill = _HEADER_FONT, _HEADER_FILL
        ws.column_dimensions[chr(64 + col)].width = width
    ws.column_dimensions["A"].hidden = True          # stable key, not for editing
    ws.freeze_panes = "B2"

    row = 2
    for key, label, guidance in LAYOUT:
        if key is None:
            ws.cell(row, 2, label).font = _SECTION_FONT
            for col in range(1, 5):
                ws.cell(row, col).fill = _SECTION_FILL
        else:
            ws.cell(row, 1, key).font = _BODY
            ws.cell(row, 2, label).font = _BODY
            v = values.get(key)
            cell = ws.cell(row, 3, v if v is not None else None)
            cell.font, cell.fill = _BODY, _INPUT_FILL
            cell.alignment = Alignment(horizontal="left")
            ws.cell(row, 4, guidance).font = _NOTE
            if key in _CHOICES:
                dv = DataValidation(type="list", formula1='"' + ",".join(_CHOICES[key]) + '"',
                                    allow_blank=False, showErrorMessage=True,
                                    errorTitle="Invalid choice",
                                    error="Pick one of: " + ", ".join(_CHOICES[key]))
                ws.add_data_validation(dv)
                dv.add(cell)
        row += 1
    _print_setup(ws, "D")

    ts = wb.create_sheet("Tickers")
    ts["A1"], ts["B1"] = "Ticker", "Notes"
    for c in (ts["A1"], ts["B1"]):
        c.font, c.fill = _HEADER_FONT, _HEADER_FILL
    ts.column_dimensions["A"].width, ts.column_dimensions["B"].width = 20, 50
    ts["B2"] = "NSE symbols, e.g. HDFCBANK (.NS is added automatically). Minimum 3."
    ts["B2"].font = _NOTE
    for i, t in enumerate(cfg.tickers, start=2):
        ts.cell(i, 1, t).font = _BODY
    _print_setup(ts, "B")

    rm = wb.create_sheet("Read me", 0)
    lines = (
        "STAR2.0 input",
        "1. Settings: change values in the yellow column only. Every setting has a default.",
        "2. Tickers: one NSE symbol per row.",
        "3. Save, then load this file in STAR2.0. Any problem is listed with its sheet and row;",
        "   nothing is silently replaced by a default.",
    )
    for i, line in enumerate(lines, start=1):
        rm.cell(i, 1, line).font = Font(name="Arial", size=12 if i == 1 else 10, bold=(i == 1))
    rm.column_dimensions["A"].width = 100

    out = Path(path)
    wb.save(out)
    return out


def _cast(key: str, raw, target: type):
    if raw is None or (isinstance(raw, str) and raw.strip() == ""):
        return None
    if target is bool:
        return str(raw).strip().lower() in ("1", "true", "yes", "y")
    if target is int:
        f = float(raw)
        if not math.isfinite(f):
            raise ValueError(f"expected a whole number, got {raw}")
        if f != int(f):
            raise ValueError(f"expected a whole number, got {raw}")
        return int(f)
    if target is float:
        text = str(raw).strip()
        value = float(text[:-1]) / 100.0 if text.endswith("%") else float(raw)
        if not math.isfinite(value):
            raise ValueError(f"expected a finite number, got {raw}")
        return value
    if key == "as_of" and hasattr(raw, "date"):
        return raw.date().isoformat()
    return str(raw).strip()


def read_input(path: str | Path) -> StarConfig:
    """Read a filled template into a validated `StarConfig`, or raise with every problem."""
    p = Path(path)
    if not p.exists():
        raise ConfigError([f"input file not found: {p}"])
    wb = openpyxl.load_workbook(p, data_only=True)
    problems: list[str] = []
    types = StarConfig.field_types()
    values: dict[str, object] = {}

    if "Settings" not in wb.sheetnames:
        problems.append("the workbook has no 'Settings' sheet (was it made from the STAR2.0 template?)")
    else:
        ws = wb["Settings"]
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, 1).value
            if not key:
                continue
            key = str(key).strip()
            if key not in types:
                problems.append(f"Settings row {r}: unknown setting key {key!r}")
                continue
            try:
                v = _cast(key, ws.cell(r, 3).value, types[key])
            except (ValueError, TypeError, OverflowError) as e:
                problems.append(f"Settings row {r} ({ws.cell(r, 2).value}): cannot read "
                                f"{ws.cell(r, 3).value!r} - {e}")
                continue
            if v is None and key not in _OPTIONAL:
                continue                         # blank = keep the default
            values[key] = v

    raw_tickers = []
    if "Tickers" in wb.sheetnames:
        ts = wb["Tickers"]
        raw_tickers = [ts.cell(r, 1).value for r in range(2, ts.max_row + 1) if ts.cell(r, 1).value]
    else:
        problems.append("the workbook has no 'Tickers' sheet")
    tickers, ticker_notes = normalise_tickers(raw_tickers)
    problems.extend(n for n in ticker_notes if not n.startswith("duplicate"))

    try:
        cfg = StarConfig.from_flat(tickers, values)
    except TypeError as e:
        raise ConfigError(problems + [f"settings could not be applied: {e}"])
    problems.extend(cfg.issues())
    if problems:
        raise ConfigError(problems)
    return cfg
