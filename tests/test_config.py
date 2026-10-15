"""Configuration: defaults, validation that reports everything, and the Excel template."""
from dataclasses import replace

import openpyxl
import pytest

from star.config.excel_input import LAYOUT, read_input, write_template
from star.config.schema import StarConfig, normalise_tickers
from star.core.errors import ConfigError

TICKERS = tuple(f"T{i}.NS" for i in range(20))
INTERNAL_KEYS = {"forecast.max_bytes"}      # engine tuning, deliberately not in the template


def test_defaults_are_valid_for_a_normal_universe():
    assert StarConfig(tickers=TICKERS).issues() == []


def test_all_problems_are_reported_together():
    cfg = StarConfig(tickers=("A.NS",))
    cfg = replace(cfg, mandate=replace(cfg.mandate, max_holdings=8),
                  backtest=replace(cfg.backtest, rebalance="weekly", window_years=9))
    problems = cfg.issues()
    joined = " | ".join(problems)
    assert "at least 3 tickers" in joined
    assert "infeasible mandate" in joined
    assert "backtest.rebalance" in joined
    assert "window_years" in joined
    with pytest.raises(ConfigError) as err:
        cfg.validated()
    assert len(err.value.issues) == len(problems)


def test_infeasible_caps_are_caught():
    """The legacy never checked this (M3-07): 8 names at 10% cannot sum to 100%."""
    cfg = StarConfig(tickers=TICKERS)
    bad = replace(cfg, mandate=replace(cfg.mandate, max_holdings=8, min_holdings=5))
    assert any("infeasible mandate" in p for p in bad.issues())


def test_ticker_normalisation():
    tickers, notes = normalise_tickers(["hdfcbank", "HDFCBANK.NS", " itc ", "", "RELIANCE.BO", "^NSEI"])
    assert tickers == ("HDFCBANK.NS", "ITC.NS", "RELIANCE.BO", "^NSEI")
    assert any("duplicate" in n for n in notes)


def test_default_horizon_is_always_in_the_grid():
    cfg = StarConfig(tickers=TICKERS)
    custom = replace(cfg, forecast=replace(cfg.forecast, horizons="1w,1y", default_horizon="45"))
    assert 45 in custom.horizon_grid().days


def test_bad_horizon_is_reported_not_raised():
    cfg = StarConfig(tickers=TICKERS)
    bad = replace(cfg, forecast=replace(cfg.forecast, horizons="1d,soon"))
    assert any("cannot read horizon" in p for p in bad.issues())


def test_template_covers_every_user_setting():
    keys = {k for k, _, _ in LAYOUT if k}
    assert set(StarConfig.field_types()) - keys == INTERNAL_KEYS


def test_template_round_trip(tmp_path):
    cfg = StarConfig(tickers=TICKERS)
    assert read_input(write_template(tmp_path / "in.xlsx", cfg)) == cfg


def test_template_is_print_ready_and_has_dropdowns(tmp_path):
    wb = openpyxl.load_workbook(write_template(tmp_path / "in.xlsx"))
    ws = wb["Settings"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert len(ws.data_validations.dataValidation) == 5


def test_bad_value_is_reported_with_its_row_not_replaced_by_a_default(tmp_path):
    path = write_template(tmp_path / "in.xlsx", StarConfig(tickers=TICKERS))
    wb = openpyxl.load_workbook(path)
    ws = wb["Settings"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "data.lookback_years":
            ws.cell(r, 3).value = "five"
            bad_row = r
    wb.save(path)
    with pytest.raises(ConfigError) as err:
        read_input(path)
    assert any(f"row {bad_row}" in p and "Lookback" in p for p in err.value.issues)


def test_percent_strings_are_understood(tmp_path):
    path = write_template(tmp_path / "in.xlsx", StarConfig(tickers=TICKERS))
    wb = openpyxl.load_workbook(path)
    ws = wb["Settings"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "mandate.max_stock_weight":
            ws.cell(r, 3).value = "8%"
    wb.save(path)
    assert read_input(path).mandate.max_stock_weight == pytest.approx(0.08)
