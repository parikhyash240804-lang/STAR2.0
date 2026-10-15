"""
Regression tests for the independent review of the P1 foundations (2026-09-17).

Each test is named after the reviewer's finding and reproduces the failing
scenario that was reported, so none of these defects can quietly return.
"""
from dataclasses import replace

import numpy as np
import pandas as pd
import pytest

from star.config.schema import StarConfig
from star.core.align import CanonicalIndex, aligned
from star.core.calendar import TradingCalendar
from star.core.errors import AlignmentError, ConfigError, ConventionError, DataError, LookAheadError
from star.core.horizons import HorizonGrid, HorizonSpec
from star.core.types import (Basis, Compounding, CovarianceEstimate, Periodicity, Provenance,
                             ReturnEstimate)
from star.data.benchmark import build_benchmark
from star.data.cache import CachedPriceProvider, DataCache
from star.data.fixture import FixtureProvider
from star.data.provider import GUARD, AsOfGuard, PricePanel
from star.data.quality import ArtifactAction, ArtifactKind, assess, repair_price_artifacts
from star.data.riskfree import RiskFreeCurve
from star.data.universe import as_of_universe
from star.data.view import DataView

TICK = ["A", "B"]


def est(values, *, basis=Basis.TOTAL, compounding=Compounding.SIMPLE, periodicity=Periodicity.ANNUAL):
    return ReturnEstimate(values=pd.Series(values, index=TICK[:len(values)]), basis=basis,
                          compounding=compounding, periodicity=periodicity,
                          provenance=Provenance("t"), periods_per_year=252.0)


def panel(n=10, tickers=("X.NS", "Y.NS")):
    idx = pd.bdate_range("2024-01-01", periods=n)
    close = pd.DataFrame({t: np.linspace(100, 110, n) for t in tickers}, index=idx)
    return PricePanel(close=close, close_raw=None, high=close * 1.01, low=close * 0.99, open_=close,
                      volume=close * 0 + 1e5, source="t", as_of=idx[-1])


# ---------------------------------------------------------------- HIGH
def test_H1_indian_day_first_dates_are_not_read_month_first(tmp_path):
    p = tmp_path / "rbi.csv"
    p.write_text("Date,Yield\n01-02-2020,5.1\n05-06-2020,3.3\n13-06-2020,3.2\n01-09-2020,3.2\n")
    c = RiskFreeCurve.from_csv(p, quote="effective")
    assert list(c.rates.index.strftime("%Y-%m-%d")) == ["2020-02-01", "2020-06-05", "2020-06-13", "2020-09-01"]
    with pytest.raises(DataError):
        c.rate("2020-01-15")          # was 3.2%, a rate first published in September


def test_H1_unreadable_dates_raise_instead_of_being_dropped(tmp_path):
    p = tmp_path / "bad.csv"
    p.write_text("Date,Yield\n2020-02-01,5.1\nsoon,3.3\n")
    with pytest.raises(DataError, match="cannot read the dates"):
        RiskFreeCurve.from_csv(p)


def test_H2_log_to_simple_needs_variance_and_gives_the_arithmetic_mean():
    mu_d, sd = 0.0005, 0.02
    log_daily = est([mu_d], compounding=Compounding.LOG, periodicity=Periodicity.DAILY)
    with pytest.raises(ConventionError):
        log_daily.as_annual_simple_total(0.0)
    via_log = log_daily.as_annual_simple_total(0.0, variance=sd ** 2).iloc[0]
    simple_daily = est([np.expm1(mu_d + 0.5 * sd ** 2)], periodicity=Periodicity.DAILY)
    via_simple = simple_daily.as_annual_simple_total(0.0).iloc[0]
    assert via_log == pytest.approx(via_simple, rel=1e-9)
    arithmetic = np.expm1(252 * mu_d + 0.5 * 252 * sd ** 2)   # exact lognormal mean, ~19.3%
    geometric = np.expm1(252 * mu_d)                           # ~13.4%: what the old code returned
    assert via_log == pytest.approx(arithmetic, rel=1e-12)
    assert via_log - geometric > 0.05


def test_H2_variance_can_come_from_a_covariance_estimate():
    cov = CovarianceEstimate(matrix=pd.DataFrame([[4e-4, 0.0], [0.0, 9e-4]], index=TICK, columns=TICK),
                             periodicity=Periodicity.DAILY, provenance=Provenance("t"))
    e = est([0.0005, 0.0003], compounding=Compounding.LOG, periodicity=Periodicity.DAILY)
    out = e.as_annual_simple_total(0.0, variance=cov)
    assert out["A"] == pytest.approx(np.expm1(0.0005 * 252 + 0.5 * 4e-4 * 252))


def test_H3_view_physically_holds_nothing_after_as_of():
    p = panel(10)
    bench = p.close["X.NS"]
    rf = pd.Series(0.065, index=p.close.index)
    view = DataView.build(p, p.close.index[4], benchmark_close=bench, risk_free=rf)
    for frame in (view.panel.close, view.panel.high, view.panel.volume, view.benchmark_close, view.risk_free):
        assert frame.index.max() <= view.as_of
    assert view.panel.as_of <= view.as_of
    view.assert_no_future_data()


def test_H3_canary_fires_if_future_data_is_smuggled_in():
    p = panel(10)
    view = DataView.build(p, p.close.index[4])
    smuggled = replace(view, panel=p)            # bypass build(): full panel inside
    with pytest.raises(LookAheadError):
        smuggled.assert_no_future_data()


# ---------------------------------------------------------------- MEDIUM
def test_M1_basis_and_compounding_conversions_commute():
    rf = 0.065
    excess = est([0.06], basis=Basis.EXCESS)                                 # annual simple excess
    a = excess.to_total(rf).values.iloc[0]
    b = excess.deannualise().to_total(rf).annualise().values.iloc[0]
    c = (excess.to_log(variance=0.0).to_total(rf).to_simple(variance=0.0).values.iloc[0])
    assert a == pytest.approx(b, abs=1e-12)
    assert a == pytest.approx(c, abs=1e-12)
    assert a == pytest.approx(1.06 * 1.065 - 1)
    assert excess.to_total(rf).to_excess(rf).values.iloc[0] == pytest.approx(0.06)


def test_M2_zero_price_is_removed_not_used_as_an_anchor():
    s = pd.Series([100.0, 101, 100, 0.0, 300, 100, 101, 102, 101, 100, 102, 103],
                  index=pd.bdate_range("2024-01-01", periods=12), name="Z.NS")
    res = repair_price_artifacts(s)
    fixed = res.close["Z.NS"]
    assert any(e.kind is ArtifactKind.NON_POSITIVE for e in res.repaired)
    assert pd.isna(fixed.iloc[3])
    assert np.isfinite(fixed.dropna()).all() and (fixed.dropna() > 0).all()
    assert fixed.iloc[4] < 150                                       # the 300 spike was repaired


def test_M3_bad_print_on_the_last_day_is_flagged_not_used_to_rescale_history():
    px = list(np.linspace(100, 105, 60)) + [210.0]                  # 2x print on the as-of day
    s = pd.Series(px, index=pd.bdate_range("2024-01-01", periods=61), name="P.NS")
    res = repair_price_artifacts(s)
    assert not res.repaired
    assert (res.close["P.NS"].values == s.values).all()


def test_M3_genuine_halving_without_volume_is_flagged_not_rescaled():
    rng = np.random.default_rng(1)
    px = 100 * np.exp(np.cumsum(rng.normal(0, 0.005, 80)))
    px[40:] *= 0.5                                                   # a real -50% crash
    s = pd.Series(px, index=pd.bdate_range("2024-01-01", periods=80), name="C.NS")
    res = repair_price_artifacts(s)
    assert not res.repaired
    assert any("no volume data" in e.note for e in res.flagged)


def test_M3_halving_confirmed_by_volume_is_repaired_as_a_split():
    rng = np.random.default_rng(2)
    idx = pd.bdate_range("2024-01-01", periods=80)
    px = 100 * np.exp(np.cumsum(rng.normal(0, 0.005, 80)))
    px[40:] *= 0.5
    vol = pd.DataFrame({"S.NS": np.where(np.arange(80) < 40, 1e5, 2e5)}, index=idx)
    res = repair_price_artifacts(pd.Series(px, index=idx, name="S.NS"), volume=vol)
    assert [e.kind for e in res.repaired] == [ArtifactKind.SPLIT_LIKE_JUMP]


def test_M4_zero_stock_cap_is_reported_not_a_crash():
    cfg = StarConfig(tickers=tuple(f"T{i}.NS" for i in range(20)))
    bad = replace(cfg, mandate=replace(cfg.mandate, max_stock_weight=0.0))
    assert any("max_stock_weight" in p for p in bad.issues())


def test_M4_infinite_cell_is_a_config_error(tmp_path):
    import openpyxl
    from star.config.excel_input import read_input, write_template
    path = write_template(tmp_path / "in.xlsx", StarConfig(tickers=tuple(f"T{i}.NS" for i in range(20))))
    wb = openpyxl.load_workbook(path)
    ws = wb["Settings"]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == "forecast.n_scenarios":
            ws.cell(r, 3).value = "inf"
    wb.save(path)
    with pytest.raises(ConfigError):
        read_input(path)


def test_M5_validation_rejects_what_it_used_to_let_through():
    cfg = StarConfig(tickers=tuple(f"T{i}.NS" for i in range(5)))
    nan_cfg = replace(cfg, data=replace(cfg.data, max_missing=float("nan")))
    assert any("finite" in p for p in nan_cfg.issues())
    assert any("infeasible mandate" in p for p in cfg.issues())      # 5 names x 10% = 50%
    assert any("YYYY-MM-DD" in p for p in replace(cfg, as_of="today").issues())
    assert any("YYYY-MM-DD" in p for p in replace(cfg, as_of="01/02/2024").issues())


def test_M6_shift_from_a_weekend_anchors_on_the_last_closed_session():
    cal = TradingCalendar.from_sessions(pd.bdate_range("2024-01-01", "2024-06-28"))
    saturday = pd.Timestamp("2024-03-30")
    assert cal.shift(saturday, 0) == pd.Timestamp("2024-03-29")
    assert cal.shift(saturday, -1) == pd.Timestamp("2024-03-28")
    assert cal.shift(saturday, 1) == pd.Timestamp("2024-04-01")


def test_M7_built_total_return_reproduces_the_proxy_over_its_calibration_window():
    p = FixtureProvider.default()
    if not (p.root / "index_NIFTYBEES_NS.parquet").exists():
        pytest.skip("reference fixture not built")
    nsei, bees = p.benchmark("^NSEI"), p.benchmark("NIFTYBEES.NS")
    b = build_benchmark(nsei, bees, as_of=nsei.index.max(), window_years=3, min_years=2)
    s, e = b.calibration_start, b.calibration_end
    window = b.tri_log_returns.loc[(b.tri_log_returns.index > s) & (b.tri_log_returns.index <= e)]
    proxy_growth = np.log(bees.loc[e] / bees.loc[s])
    # exact when the two calendars coincide; within a few bp when they differ by a session or two
    assert window.sum() == pytest.approx(proxy_growth, abs=0.002)


# ---------------------------------------------------------------- LOW
def test_L1_zero_lookback_is_rejected():
    p = panel(10)
    view = DataView.build(p, p.close.index[-1])
    with pytest.raises(DataError):
        view.closes(lookback_sessions=0)
    with pytest.raises(DataError):
        as_of_universe(p.close, p.close.index[-1], lookback_sessions=0, min_sessions=1, min_adv_inr=0, min_names=1)


def test_L2_date_indexed_rf_series_is_refused():
    e = est([0.06, 0.07])
    with pytest.raises(ConventionError):
        e.to_total(pd.Series([0.065], index=pd.DatetimeIndex(["2024-01-01"])))


def test_L3_covariance_with_scrambled_rows_is_fully_aligned():
    canon = CanonicalIndex.of(["A", "B"])
    m = pd.DataFrame([[2.0, 1.0], [1.0, 3.0]], index=["B", "A"], columns=["A", "B"])  # rows B,A
    out = canon.align(m)
    assert list(out.index) == ["A", "B"] and list(out.columns) == ["A", "B"]
    assert out.loc["A", "A"] == 1.0 and out.loc["B", "B"] == 1.0       # values follow their labels


def test_L3_aligned_decorator_aligns_typed_estimates():
    @aligned("mu")
    def first(index: CanonicalIndex, mu):
        return list(mu.values.index)
    e = ReturnEstimate(values=pd.Series([1.0, 2.0], index=["B", "A"]), basis=Basis.TOTAL,
                       compounding=Compounding.SIMPLE, periodicity=Periodicity.ANNUAL,
                       provenance=Provenance("t"))
    assert first(CanonicalIndex.of(["A", "B"]), e) == ["A", "B"]


def test_L4_weights_for_unknown_tickers_raise():
    cov = CovarianceEstimate(matrix=pd.DataFrame(np.eye(2), index=TICK, columns=TICK),
                             periodicity=Periodicity.ANNUAL, provenance=Provenance("t"))
    with pytest.raises(AlignmentError):
        cov.portfolio_variance(pd.Series({"A": 0.5, "C": 0.5}))


def test_L5_sub_day_horizons_raise_and_custom_keeps_labels():
    for token in ("0", "0.2d"):
        with pytest.raises(ConfigError):
            HorizonSpec.parse(token)
    g = HorizonGrid.parse("1m,3m,1y", periods_per_year=248).with_custom(45)
    assert "3M" in g.labels and "1Y" in g.labels and 45 in g.days


def test_L6_short_history_annualisation_is_not_inflated():
    cal = TradingCalendar.from_sessions(pd.bdate_range("2024-01-01", periods=130))
    assert cal.periods_per_year == pytest.approx(261, abs=3)
    with pytest.raises(DataError):
        TradingCalendar.from_sessions(pd.bdate_range("2024-01-01", periods=20)).periods_per_year


def test_L7_ambiguous_units_are_recorded():
    c = RiskFreeCurve.from_series(pd.Series([0.14, 0.08], index=pd.to_datetime(["2021-01-01", "2021-02-01"])),
                                  name="x", source="x")
    assert c.degradations and "ambiguous" in c.degradations[0]


def test_L8_tbill_simple_quote_is_converted_to_effective():
    c = RiskFreeCurve.from_series(pd.Series([6.5], index=pd.to_datetime(["2024-01-01"])),
                                  name="tbill", source="t", quote="tbill_simple")
    assert c.rate("2024-01-02") == pytest.approx((1 + 0.065 * 91 / 365) ** (365 / 91) - 1)
    assert c.rate("2024-01-02") == pytest.approx(0.0666, abs=2e-4)


def test_L9_cache_replays_in_requested_order_and_caps_as_of(tmp_path):
    class Inner:
        name = "inner"
        def prices(self, tickers, start, end, *, adjusted=True):
            return panel(30, tickers=("A.NS", "B.NS"))
    cp = CachedPriceProvider(Inner(), DataCache(tmp_path))
    assert list(cp.prices(["B.NS", "A.NS"], "2024-01-01", "2024-02-09").close.columns) == ["B.NS", "A.NS"]
    with GUARD.armed("2024-01-10"):
        got = cp.prices(["A.NS", "B.NS"], "2024-01-01", "2024-01-10")
    assert got.as_of <= pd.Timestamp("2024-01-10")


def test_L10_open_ended_request_while_armed_raises_and_tz_index_clamps():
    guard = AsOfGuard()
    with guard.armed("2024-01-05"):
        with pytest.raises(LookAheadError):
            guard.check(None, what="t")
        idx = pd.date_range("2024-01-01", periods=10, freq="D", tz="Asia/Kolkata")
        out = guard.clamp(pd.Series(range(10), index=idx), what="t")
        assert len(out) == 5


def test_L11_missing_volume_does_not_pass_the_liquidity_floor():
    idx = pd.bdate_range("2024-01-01", periods=300)
    close = pd.DataFrame({"A.NS": 100.0, "B.NS": 100.0}, index=idx)
    volume = pd.DataFrame({"A.NS": 1e6}, index=idx)            # no volume at all for B
    rep = assess(close, volume=volume, min_adv_inr=1e6, min_names=1)
    b = next(v for v in rep.verdicts if v.ticker == "B.NS")
    assert b.dropped and "no volume" in b.reason


def test_L12_fractional_calibration_window_works():
    p = FixtureProvider.default()
    if not (p.root / "index_NIFTYBEES_NS.parquet").exists():
        pytest.skip("reference fixture not built")
    b = build_benchmark(p.benchmark("^NSEI"), p.benchmark("NIFTYBEES.NS"),
                        as_of="2026-09-11", window_years=2.5, min_years=2)
    assert b.calibrated
